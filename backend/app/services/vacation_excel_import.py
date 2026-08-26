"""Парсинг графика отпусков (форма Т-7 и похожие .xlsx / .xlsm)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

_FIO_HEADERS = frozenset({"фамилияимяотчество", "фио", "фиосотрудника"})
_TAB_HEADERS = frozenset({"табельныйномер", "табельный", "табномер"})
_START_HEADERS = frozenset({"датаначалаотпуска", "началоотпуска", "датаначала"})
_END_HEADERS = frozenset({"датаокончанияотпуска", "окончаниеотпуска", "датаокончания", "датаконцаотпуска"})
_KIND_HEADERS = frozenset({"видотпуска", "типотпуска"})
_SKIP_NAME = re.compile(
    r"^(январь|февраль|март|апрель|май|июнь|июль|август|сентябрь|октябрь|ноябрь|декабрь|"
    r"всегоза|итого)",
    re.IGNORECASE,
)


def fold_header(val: object) -> str:
    if val is None:
        return ""
    s = str(val).strip().casefold().replace("ё", "е")
    return re.sub(r"[\s.\-_/(),]+", "", s)


def fold_name(val: str) -> str:
    return re.sub(r"\s+", " ", (val or "").replace("ё", "е").replace("Ё", "е")).strip().casefold()


def fold_personnel(val: str) -> str:
    digits = re.sub(r"\D", "", val or "")
    return digits.lstrip("0")


def cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() or abs(value - round(value)) < 1e-6:
            return str(int(round(value)))
        return str(value).strip()
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value: object) -> date | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        serial = int(round(float(value)))
        if 20000 <= serial <= 80000:
            return date(1899, 12, 30) + timedelta(days=serial)
        return None
    text = cell_text(value)
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def vacation_kind(raw: str | None) -> str:
    t = fold_header(raw or "")
    if "учебн" in t:
        return "study"
    if "больнич" in t or t in {"б", "бл"}:
        return "sick"
    return "vacation"


@dataclass(frozen=True)
class ParsedVacationRow:
    sheet_row: int
    full_name: str | None
    personnel_number: str | None
    start: date | None
    end: date | None
    kind: str
    error: str | None = None


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    max_scan = min(ws.max_row or 0, 40)
    max_col = min(ws.max_column or 0, 40)
    for r in range(1, max_scan + 1):
        for c in range(1, max_col + 1):
            key = fold_header(ws.cell(r, c).value)
            if not key:
                continue
            if key in _FIO_HEADERS and "name" not in mapping:
                mapping["name"] = c
            elif key in _TAB_HEADERS and "tab" not in mapping:
                mapping["tab"] = c
            elif key in _START_HEADERS and "start" not in mapping:
                mapping["start"] = c
            elif key in _END_HEADERS and "end" not in mapping:
                mapping["end"] = c
            elif key in _KIND_HEADERS and "kind" not in mapping:
                mapping["kind"] = c
    if "name" in mapping and "start" not in mapping:
        # Типовая Т-7: ФИО=D, табельный=E, вид=F, начало=H, конец=I
        mapping.setdefault("tab", mapping["name"] + 1)
        mapping.setdefault("kind", mapping["name"] + 2)
        mapping.setdefault("start", mapping["name"] + 4)
        mapping.setdefault("end", mapping["name"] + 5)
    return mapping


def _is_section_row(name: str | None, start: date | None, end: date | None) -> bool:
    if start or end:
        return False
    if not name:
        return True
    compact = fold_name(name)
    if _SKIP_NAME.match(compact.replace(" ", "")):
        return True
    if "филиал" in compact and len(name) > 40:
        return True
    if compact.startswith("всего "):
        return True
    return False


def parse_vacation_excel(content: bytes) -> tuple[list[ParsedVacationRow], str | None]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as e:  # noqa: BLE001
        return [], f"Не удалось открыть файл как Excel (.xlsx/.xlsm): {e!s}"
    try:
        if not wb.sheetnames:
            return [], "В файле нет листов"
        ws = wb[wb.sheetnames[0]]
        cols = _header_map(ws)
        if "name" not in cols and "tab" not in cols:
            return [], "Не найдены колонки «ФИО» или «Табельный номер»"
        if "start" not in cols or "end" not in cols:
            return [], "Не найдены колонки даты начала и окончания отпуска"
        name_c = cols.get("name")
        tab_c = cols.get("tab")
        start_c = cols["start"]
        end_c = cols["end"]
        kind_c = cols.get("kind")
        rows: list[ParsedVacationRow] = []
        max_row = ws.max_row or 0
        for r in range(1, max_row + 1):
            name = cell_text(ws.cell(r, name_c).value) if name_c else ""
            tab = cell_text(ws.cell(r, tab_c).value) if tab_c else ""
            start = parse_date(ws.cell(r, start_c).value)
            end = parse_date(ws.cell(r, end_c).value)
            kind_raw = cell_text(ws.cell(r, kind_c).value) if kind_c else ""
            if fold_header(name) in _FIO_HEADERS or fold_header(name) == "4":
                continue
            if _is_section_row(name or None, start, end) and not fold_personnel(tab):
                continue
            if not name and not tab and not start and not end:
                continue
            if not start or not end:
                if len(fold_personnel(tab)) >= 4:
                    rows.append(
                        ParsedVacationRow(
                            sheet_row=r,
                            full_name=name or None,
                            personnel_number=tab or None,
                            start=start,
                            end=end,
                            kind=vacation_kind(kind_raw),
                            error="Нет даты начала или окончания",
                        )
                    )
                continue
            if end < start:
                rows.append(
                    ParsedVacationRow(
                        sheet_row=r,
                        full_name=name or None,
                        personnel_number=tab or None,
                        start=start,
                        end=end,
                        kind=vacation_kind(kind_raw),
                        error="Дата окончания раньше начала",
                    )
                )
                continue
            rows.append(
                ParsedVacationRow(
                    sheet_row=r,
                    full_name=name or None,
                    personnel_number=tab or None,
                    start=start,
                    end=end,
                    kind=vacation_kind(kind_raw),
                )
            )
        if not rows:
            return [], "В файле нет строк с датами отпуска"
        return rows, None
    finally:
        wb.close()
