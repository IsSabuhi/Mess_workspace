"""Parse Excel of GSM SIM cards: phone, ICCID, IP, installation address."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from openpyxl import load_workbook

# H, K, Q, R — как в шаблоне; заголовки в 1-й строке могут сдвинуть колонки.
_COL_PHONE = 8
_COL_ICCID = 11
_COL_IP = 17
_COL_ADDR = 18

_PHONE_HEADERS = frozenset({"номертелефона", "телефон", "phone", "msisdn"})
_ICCID_HEADERS = frozenset(
    {"номерзасимкарты", "номерзаsimкарты", "iccid", "sim", "simкарта", "номерsim"}
)
_IP_HEADERS = frozenset({"ip", "ipадрес", "ip-адрес", "адресip"})
_ADDR_HEADERS = frozenset({"адресустановки", "адрес", "address", "location"})

IPV4_RE = re.compile(
    r"\b(?:(?:25[0-5]|2[0-4]\d|[01]?\d\d?)\.){3}(?:25[0-5]|2[0-4]\d|[01]?\d\d?)\b"
)


@dataclass(frozen=True)
class ParsedSimRow:
    sheet_row: int
    phone: str | None
    iccid: str | None
    ip: str | None
    address: str | None


def extract_ipv4(text: str | None) -> set[str]:
    if not text:
        return set()
    return set(IPV4_RE.findall(text.strip()))


def cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() or abs(value - round(value)) < 1e-6:
            return str(int(round(value)))
        return str(value).strip()
    return re.sub(r"\s+", " ", str(value)).strip()


def _norm_header(val: object) -> str:
    if val is None:
        return ""
    s = str(val).strip().lower().replace("ё", "е")
    return re.sub(r"[\s.\-_/]+", "", s)


def _header_map(ws) -> dict[str, int]:
    mapping: dict[str, int] = {}
    max_col = min(ws.max_column or 0, 40)
    for c in range(1, max_col + 1):
        key = _norm_header(ws.cell(1, c).value)
        if not key:
            continue
        if key in _PHONE_HEADERS and "phone" not in mapping:
            mapping["phone"] = c
        elif (key in _ICCID_HEADERS or "sim" in key and "номер" in key) and "iccid" not in mapping:
            mapping["iccid"] = c
        elif key in _IP_HEADERS and "ip" not in mapping:
            mapping["ip"] = c
        elif key in _ADDR_HEADERS and "address" not in mapping:
            mapping["address"] = c
    return mapping


def parse_uspd_sim_excel(content: bytes) -> tuple[list[ParsedSimRow], str | None]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as e:  # noqa: BLE001
        return [], f"Не удалось открыть файл как Excel (.xlsx): {e!s}"
    try:
        if not wb.sheetnames:
            return [], "В файле нет листов"
        ws = wb[wb.sheetnames[0]]
        cols = _header_map(ws)
        phone_c = cols.get("phone", _COL_PHONE)
        iccid_c = cols.get("iccid", _COL_ICCID)
        ip_c = cols.get("ip", _COL_IP)
        addr_c = cols.get("address", _COL_ADDR)
        max_row = ws.max_row or 0
        if max_row < 2:
            return [], "Нет строк данных после заголовка"
        rows: list[ParsedSimRow] = []
        for r in range(2, max_row + 1):
            phone = cell_text(ws.cell(r, phone_c).value) or None
            iccid = cell_text(ws.cell(r, iccid_c).value) or None
            ip = cell_text(ws.cell(r, ip_c).value) or None
            address = cell_text(ws.cell(r, addr_c).value) or None
            if not any((phone, iccid, ip, address)):
                continue
            if iccid and len(iccid) > 32:
                iccid = iccid[:32]
            if phone and len(phone) > 64:
                phone = phone[:64]
            rows.append(
                ParsedSimRow(sheet_row=r, phone=phone, iccid=iccid, ip=ip, address=address)
            )
        if not rows:
            return [], "Нет заполненных строк (нужны телефон, SIM, IP или адрес)"
        return rows, None
    finally:
        wb.close()


def is_gsm_object(name: str | None) -> bool:
    t = (name or "").strip().lower()
    if not t:
        return False
    if t == "gsm" or t.startswith("gsm "):
        return True
    return any(x in t for x in ("телефон", "телеофис", "телефонис"))
